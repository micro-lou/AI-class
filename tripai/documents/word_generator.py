from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from typing import Dict, Any, List
import os
from utils.helpers import format_currency

class WordGenerator:
    def __init__(self):
        self.doc = Document()
    
    def _add_title(self, text: str, level: int = 1):
        heading = self.doc.add_heading(text, level=level)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if level == 1:
            heading.runs[0].font.size = Pt(24)
            heading.runs[0].font.bold = True
            heading.runs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    def _add_paragraph(self, text: str, bold: bool = False, font_size: int = 11):
        para = self.doc.add_paragraph(text)
        para.runs[0].font.bold = bold
        para.runs[0].font.size = Pt(font_size)
        return para
    
    def _add_table(self, data: List[List[str]], has_header: bool = True):
        rows, cols = len(data), len(data[0]) if data else 0
        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for i, row in enumerate(data):
            for j, cell in enumerate(row):
                table.cell(i, j).text = str(cell)
                table.cell(i, j).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if has_header and i == 0:
                    table.cell(i, j).paragraphs[0].runs[0].font.bold = True
                    table.cell(i, j).paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    def _add_section_break(self):
        self.doc.add_page_break()
    
    def generate_cover(self, itinerary: Dict[str, Any]):
        title = self.doc.add_heading("智能旅行攻略", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].font.size = Pt(32)
        title.runs[0].font.bold = True
        
        self.doc.add_paragraph()
        route_title = self.doc.add_heading(
            f"{itinerary.get('origin', '')} → {itinerary['destination']}", level=1
        )
        route_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        sf = itinerary.get('selected_flight')
        srf = itinerary.get('selected_return_flight')
        has_flight = sf and sf is not None
        has_return_flight = srf and srf is not None
        row_count = 6
        if itinerary.get('departure_date'):
            row_count += 1
        if has_flight:
            row_count += 2
        if has_return_flight:
            row_count += 2
        
        info_table = self.doc.add_table(rows=row_count, cols=2)
        info_table.style = "Table Grid"
        info_data = [
            ["旅行天数", f"{itinerary['days']}天"],
            ["出行人群", itinerary['travel_group']],
            ["游玩人数", f"{itinerary.get('traveler_count', 1)}人"],
            ["往返交通", itinerary.get('transport', '')],
            ["通勤费用", format_currency(itinerary.get('transport_cost', 0))],
            ["总预算（含通勤）", format_currency(itinerary['total_budget'])]
        ]
        
        if itinerary.get('departure_date'):
            info_data.append(["出行日期", f"{itinerary.get('departure_date', '')} ~ {itinerary.get('return_date', '')}"])
        
        if has_flight:
            transport = itinerary.get('transport', '')
            flight_label = "航班" if transport == "飞机" else "车次"
            info_data.append([
                f"去程{flight_label}",
                f"{sf.get('flight_no', '')} | {sf.get('departure_time', '')}→{sf.get('arrival_time', '')} | {sf.get('duration', '')} | ¥{sf.get('price_per_person', 0)}/人"
            ])
            info_data.append([
                f"{'航空公司' if transport == '飞机' else '运营方'}",
                f"{sf.get('airline_or_operator', '')} · {'舱位' if transport == '飞机' else '座席'}：{sf.get('seat_type', '')}"
            ])
        
        if has_return_flight:
            transport = itinerary.get('transport', '')
            flight_label = "航班" if transport == "飞机" else "车次"
            info_data.append([
                f"返程{flight_label}",
                f"{srf.get('flight_no', '')} | {srf.get('departure_time', '')}→{srf.get('arrival_time', '')} | {srf.get('duration', '')} | ¥{srf.get('price_per_person', 0)}/人"
            ])
            info_data.append([
                f"返程{'航空公司' if transport == '飞机' else '运营方'}",
                f"{srf.get('airline_or_operator', '')} · {'舱位' if transport == '飞机' else '座席'}：{srf.get('seat_type', '')}"
            ])
        
        for i, row in enumerate(info_data):
            for j, cell in enumerate(row):
                info_table.cell(i, j).text = cell
                if j == 0:
                    info_table.cell(i, j).paragraphs[0].runs[0].font.bold = True
        
        summary_para = self.doc.add_paragraph()
        summary_para.add_run("行程概述：").bold = True
        summary_para.add_run(
            itinerary['summary'][:80] + "..." if len(itinerary['summary']) > 80 else itinerary['summary']
        )
        
        self._add_section_break()
    
    def generate_toc(self):
        self._add_title("目录")
        toc_items = ["行程概述", "每日行程", "美食推荐", "购物指南", "预算明细", "注意事项", "出行清单"]
        for item in toc_items:
            para = self.doc.add_paragraph()
            para.add_run(item).font.size = Pt(12)
        self._add_section_break()
    
    def generate_summary(self, itinerary: Dict[str, Any]):
        self._add_title("行程概述")
        self._add_paragraph(itinerary['summary'])
        
        sf = itinerary.get('selected_flight')
        srf = itinerary.get('selected_return_flight')
        has_flight = sf and sf is not None
        has_return_flight = srf and srf is not None
        base_rows = 6
        if itinerary.get('departure_date'):
            base_rows += 1
        if has_flight:
            base_rows += 1
        if has_return_flight:
            base_rows += 1
        
        stats_table = self.doc.add_table(rows=base_rows, cols=2)
        stats_table.style = "Table Grid"
        local_cost = itinerary['total_budget'] - itinerary.get('transport_cost', 0)
        stats_data = [
            ["出发地", itinerary.get('origin', '')],
            ["目的地", itinerary['destination']],
            ["总天数", f"{itinerary['days']}天{itinerary['days']-1}晚"],
            ["游玩人数", f"{itinerary.get('traveler_count', 1)}人 · {itinerary['travel_group']}"],
            ["往返交通", f"{itinerary.get('transport', '')} · 通勤费用 {format_currency(itinerary.get('transport_cost', 0))}"],
            ["预算分配", f"通勤 {format_currency(itinerary.get('transport_cost', 0))} + 当地 {format_currency(local_cost)} = {format_currency(itinerary['total_budget'])}"]
        ]
        
        if itinerary.get('departure_date'):
            stats_data.append(["出行日期", f"{itinerary.get('departure_date', '')} ~ {itinerary.get('return_date', '')}"])
        
        if has_flight:
            transport = itinerary.get('transport', '')
            flight_label = "航班" if transport == "飞机" else "车次"
            stats_data.append([
                f"去程{flight_label}",
                f"{sf.get('flight_no', '')} {sf.get('departure_time', '')}→{sf.get('arrival_time', '')} 抵达后安排首日行程"
            ])
        
        if has_return_flight:
            transport = itinerary.get('transport', '')
            flight_label = "航班" if transport == "飞机" else "车次"
            stats_data.append([
                f"返程{flight_label}",
                f"{srf.get('flight_no', '')} {srf.get('departure_time', '')}→{srf.get('arrival_time', '')} 出发前安排末尾天行程"
            ])
        
        for i, row in enumerate(stats_data):
            for j, cell in enumerate(row):
                stats_table.cell(i, j).text = cell
                if j == 0:
                    stats_table.cell(i, j).paragraphs[0].runs[0].font.bold = True
        
        self._add_section_break()
    
    def generate_daily_itinerary(self, itinerary: Dict[str, Any]):
        self._add_title("每日行程")
        
        for day in itinerary['daily_itinerary']:
            self._add_title(f"第{day['day']}天 - {day['title']}", 2)
            
            self._add_paragraph("上午行程", bold=True)
            self._add_paragraph(f"景点：{day['morning']['attraction']}")
            self._add_paragraph(f"介绍：{day['morning']['description']}")
            self._add_paragraph(f"时长：{day['morning']['duration']}")
            self._add_paragraph(f"门票：{format_currency(day['morning']['ticket_price'])}")
            
            self._add_paragraph("午餐推荐", bold=True)
            self._add_paragraph(f"餐厅：{day['lunch']['restaurant']}")
            self._add_paragraph(f"推荐菜：{day['lunch']['dish']}")
            
            self._add_paragraph("下午行程", bold=True)
            self._add_paragraph(f"景点：{day['afternoon']['attraction']}")
            self._add_paragraph(f"介绍：{day['afternoon']['description']}")
            
            self._add_paragraph("晚餐推荐", bold=True)
            self._add_paragraph(f"餐厅：{day['dinner']['restaurant']}")
            
            self._add_paragraph("住宿安排", bold=True)
            acc = day['accommodation']
            self._add_paragraph(f"酒店：{acc.get('hotel', '')}")
            self._add_paragraph(f"位置：{acc.get('location', '')}")
            if acc.get('price_per_night'):
                self._add_paragraph(f"每晚价格：{format_currency(acc['price_per_night'])}")
            if acc.get('nearby_attractions'):
                self._add_paragraph(f"附近景点：{', '.join(acc['nearby_attractions'])}")
            if acc.get('selection_reason'):
                self._add_paragraph(f"选择理由：{acc['selection_reason']}")
            
            commutes = day.get('commutes', [])
            if commutes:
                self._add_paragraph("景点间通勤", bold=True)
                commute_data = [["出发", "到达", "方式", "耗时", "费用/人", "路线详情"]]
                for c in commutes:
                    commute_data.append([
                        c.get('from', ''),
                        c.get('to', ''),
                        c.get('method', ''),
                        c.get('duration', ''),
                        format_currency(c.get('cost_per_person', 0)),
                        c.get('detail', '')
                    ])
                self._add_table(commute_data)
            
            self._add_paragraph(f"当日预算：{format_currency(day['daily_budget'])}", bold=True)
            self.doc.add_paragraph()
        
        self._add_section_break()
    
    def generate_food_recommendations(self, itinerary: Dict[str, Any]):
        self._add_title("美食推荐")
        food_data = [["美食名称", "介绍", "价格", "推荐地点"]]
        for food in itinerary['food_recommendations']:
            food_data.append([food['name'], food['description'], format_currency(food['price']), food['location']])
        self._add_table(food_data)
        self._add_section_break()
    
    def generate_shopping(self, itinerary: Dict[str, Any]):
        self._add_title("购物指南")
        shop_data = [["购物地点", "介绍", "特色商品"]]
        for shop in itinerary['shopping_recommendations']:
            shop_data.append([shop['name'], shop['description'], shop['specialty']])
        self._add_table(shop_data)
        self._add_section_break()
    
    def generate_budget(self, itinerary: Dict[str, Any]):
        self._add_title("预算明细")
        budget_data = [["项目", "金额"]]
        budget_data.append([f"往返交通（{itinerary.get('transport', '')}，{itinerary.get('traveler_count', 1)}人）", format_currency(itinerary.get('transport_cost', 0))])
        for day in itinerary['daily_itinerary']:
            budget_data.append([f"第{day['day']}天", format_currency(day['daily_budget'])])
        budget_data.append(["总计", format_currency(itinerary['total_budget'])])
        self._add_table(budget_data)
        self._add_section_break()
    
    def generate_tips(self, itinerary: Dict[str, Any]):
        self._add_title("注意事项")
        for i, tip in enumerate(itinerary['tips'], 1):
            self._add_paragraph(f"{i}. {tip}")
        self._add_section_break()
    
    def generate_packing_list(self, itinerary: Dict[str, Any]):
        self._add_title("出行清单")
        for item in itinerary['packing_list']:
            self._add_paragraph(f"• {item}")
    
    def generate_document(self, itinerary: Dict[str, Any], output_path: str):
        self.generate_cover(itinerary)
        self.generate_toc()
        self.generate_summary(itinerary)
        self.generate_daily_itinerary(itinerary)
        self.generate_food_recommendations(itinerary)
        self.generate_shopping(itinerary)
        self.generate_budget(itinerary)
        self.generate_tips(itinerary)
        self.generate_packing_list(itinerary)
        self.doc.save(output_path)
        return output_path
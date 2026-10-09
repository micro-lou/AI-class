from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from typing import Dict, Any, List
import os
from utils.helpers import format_currency

class PPTGenerator:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(16)
        self.prs.slide_height = Inches(9)
    
    def _add_title_slide(self, title: str, subtitle: str = ""):
        slide_layout = self.prs.slide_layouts[0]
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_holder = slide.shapes.title
        title_holder.text = title
        title_holder.text_frame.paragraphs[0].font.size = Pt(44)
        title_holder.text_frame.paragraphs[0].font.bold = True
        
        if subtitle:
            subtitle_holder = slide.placeholders[1]
            subtitle_holder.text = subtitle
    
    def _add_title_only_slide(self, title: str):
        slide_layout = self.prs.slide_layouts[5]
        slide = self.prs.slides.add_slide(slide_layout)
        title_holder = slide.shapes.title
        title_holder.text = title
        title_holder.text_frame.paragraphs[0].font.size = Pt(36)
        title_holder.text_frame.paragraphs[0].font.bold = True
    
    def _add_bullet_points(self, slide, points: List[str], left: int, top: int, width: int, height: int):
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txBox.text_frame
        
        for i, point in enumerate(points):
            p = tf.add_paragraph()
            p.text = point
            p.font.size = Pt(16)
            if i > 0:
                p.level = 1
    
    def generate_cover(self, itinerary: Dict[str, Any]):
        title = f"{itinerary.get('origin', '')} → {itinerary['destination']} 旅行攻略"
        sf = itinerary.get('selected_flight')
        srf = itinerary.get('selected_return_flight')
        flight_info = ""
        if sf and sf is not None:
            flight_label = "航班" if itinerary.get('transport') == "飞机" else "车次"
            flight_info = f" | 去程{sf.get('flight_no', '')} {sf.get('departure_time', '')}→{sf.get('arrival_time', '')}"
        if srf and srf is not None:
            flight_label = "航班" if itinerary.get('transport') == "飞机" else "车次"
            flight_info += f" | 返程{srf.get('flight_no', '')} {srf.get('departure_time', '')}→{srf.get('arrival_time', '')}"
        date_info = ""
        if itinerary.get('departure_date'):
            date_info = f" | {itinerary.get('departure_date', '')}~{itinerary.get('return_date', '')}"
        subtitle = (
            f"{itinerary['days']}天行程 · {itinerary['travel_group']} · "
            f"{itinerary.get('traveler_count', 1)}人 · {itinerary.get('transport', '')} · "
            f"预算{format_currency(itinerary['total_budget'])}（含通勤{format_currency(itinerary.get('transport_cost', 0))}）"
            f"{date_info}{flight_info}"
        )
        self._add_title_slide(title, subtitle)
    
    def generate_overview(self, itinerary: Dict[str, Any]):
        slide_layout = self.prs.slide_layouts[5]
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_holder = slide.shapes.title
        title_holder.text = "行程概览"
        
        local_cost = itinerary['total_budget'] - itinerary.get('transport_cost', 0)
        overview_points = [
            f"出发地：{itinerary.get('origin', '')}",
            f"目的地：{itinerary['destination']}",
            f"旅行天数：{itinerary['days']}天{itinerary['days']-1}晚",
            f"游玩人数：{itinerary.get('traveler_count', 1)}人",
            f"出行人群：{itinerary['travel_group']}",
            f"往返交通：{itinerary.get('transport', '')}（通勤费用 {format_currency(itinerary.get('transport_cost', 0))}）",
            f"总预算：{format_currency(itinerary['total_budget'])}（通勤 {format_currency(itinerary.get('transport_cost', 0))} + 当地 {format_currency(local_cost)}）",
        ]
        
        if itinerary.get('departure_date'):
            overview_points.append(f"出行日期：{itinerary.get('departure_date', '')} ~ {itinerary.get('return_date', '')}")
        
        sf = itinerary.get('selected_flight')
        if sf and sf is not None:
            transport = itinerary.get('transport', '')
            flight_label = "航班" if transport == "飞机" else "车次"
            overview_points.append("")
            overview_points.append(f"去程{flight_label}：{sf.get('flight_no', '')}")
            overview_points.append(f"   出发 {sf.get('departure_time', '')} → 抵达 {sf.get('arrival_time', '')}（{sf.get('duration', '')}）")
            overview_points.append(f"   {'航空公司' if transport == '飞机' else '运营方'}：{sf.get('airline_or_operator', '')} · {'舱位' if transport == '飞机' else '座席'}：{sf.get('seat_type', '')}")
            overview_points.append(f"   单人 ¥{sf.get('price_per_person', 0)} | {itinerary.get('traveler_count', 1)}人总价 ¥{sf.get('price_total', 0)}")
            overview_points.append(f"   ⚠️ 第1天行程将按抵达时间 {sf.get('arrival_time', '')} 安排")
        
        srf = itinerary.get('selected_return_flight')
        if srf and srf is not None:
            transport = itinerary.get('transport', '')
            flight_label = "航班" if transport == "飞机" else "车次"
            overview_points.append("")
            overview_points.append(f"返程{flight_label}：{srf.get('flight_no', '')}")
            overview_points.append(f"   出发 {srf.get('departure_time', '')} → 抵达 {srf.get('arrival_time', '')}（{srf.get('duration', '')}）")
            overview_points.append(f"   {'航空公司' if transport == '飞机' else '运营方'}：{srf.get('airline_or_operator', '')} · {'舱位' if transport == '飞机' else '座席'}：{srf.get('seat_type', '')}")
            overview_points.append(f"   单人 ¥{srf.get('price_per_person', 0)} | {itinerary.get('traveler_count', 1)}人总价 ¥{srf.get('price_total', 0)}")
            overview_points.append(f"   ⚠️ 最后1天行程将按出发时间 {srf.get('departure_time', '')} 安排")
        
        overview_points.append("")
        overview_points.append(f"行程简介：")
        overview_points.append(f"{itinerary['summary']}")
        
        self._add_bullet_points(slide, overview_points, 1, 2, 14, 6)
    
    def generate_daily_slides(self, itinerary: Dict[str, Any]):
        for day in itinerary['daily_itinerary']:
            slide_layout = self.prs.slide_layouts[5]
            slide = self.prs.slides.add_slide(slide_layout)
            
            title_holder = slide.shapes.title
            title_holder.text = f"第{day['day']}天 - {day['title']}"
            
            day_points = [
                f"上午：{day['morning']['attraction']}",
                f"   - {day['morning']['description'][:50]}...",
                f"   - 时长：{day['morning']['duration']} | 门票：{format_currency(day['morning']['ticket_price'])}",
                f"",
                f"午餐：{day['lunch']['restaurant']}",
                f"   - 推荐：{day['lunch']['dish']} | 人均：{format_currency(day['lunch']['price_per_person'])}",
                f"",
                f"下午：{day['afternoon']['attraction']}",
                f"   - {day['afternoon']['description'][:50]}...",
                f"   - 时长：{day['afternoon']['duration']} | 门票：{format_currency(day['afternoon']['ticket_price'])}",
                f"",
                f"晚餐：{day['dinner']['restaurant']}",
                f"   - 推荐：{day['dinner']['dish']} | 人均：{format_currency(day['dinner']['price_per_person'])}",
                f"",
                f"住宿：{day['accommodation'].get('hotel', '')}",
                f"   - 位置：{day['accommodation'].get('location', '')} | 价格：{format_currency(day['accommodation'].get('price_per_night', 0))}",
            ]
            
            nearby = day['accommodation'].get('nearby_attractions', [])
            if nearby:
                day_points.append(f"   - 附近景点：{', '.join(nearby)}")
            reason = day['accommodation'].get('selection_reason', '')
            if reason:
                day_points.append(f"   - 选择理由：{reason[:40]}{'...' if len(reason) > 40 else ''}")
            
            day_points.append("")
            
            commutes = day.get('commutes', [])
            if commutes:
                day_points.append("景点间通勤：")
                for c in commutes:
                    day_points.append(
                        f"   {c.get('from', '')} → {c.get('to', '')}："
                        f"{c.get('method', '')} {c.get('duration', '')} "
                        f"{format_currency(c.get('cost_per_person', 0))}/人"
                    )
                day_points.append("")
            
            day_points.append(f"当日预算：{format_currency(day['daily_budget'])}")
            
            self._add_bullet_points(slide, day_points, 1, 2, 14, 6)
    
    def generate_food_slide(self, itinerary: Dict[str, Any]):
        slide_layout = self.prs.slide_layouts[5]
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_holder = slide.shapes.title
        title_holder.text = "美食推荐"
        
        food_points = []
        for food in itinerary['food_recommendations']:
            food_points.append(f"{food['name']}")
            food_points.append(f"   - {food['description'][:40]}...")
            food_points.append(f"   - 价格：{format_currency(food['price'])} | 地点：{food['location']}")
            food_points.append("")
        
        self._add_bullet_points(slide, food_points, 1, 2, 14, 6)
    
    def generate_budget_slide(self, itinerary: Dict[str, Any]):
        slide_layout = self.prs.slide_layouts[5]
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_holder = slide.shapes.title
        title_holder.text = "预算明细"
        
        budget_points = [f"往返交通（{itinerary.get('transport', '')}，{itinerary.get('traveler_count', 1)}人）："]
        budget_points.append(f"   {format_currency(itinerary.get('transport_cost', 0))}")
        budget_points.append("")
        budget_points.append("每日预算：")
        for day in itinerary['daily_itinerary']:
            budget_points.append(f"   第{day['day']}天：{format_currency(day['daily_budget'])}")
        
        budget_points.append("")
        budget_points.append(f"总预算（含通勤）：{format_currency(itinerary['total_budget'])}")
        
        self._add_bullet_points(slide, budget_points, 1, 2, 14, 6)
    
    def generate_tips_slide(self, itinerary: Dict[str, Any]):
        slide_layout = self.prs.slide_layouts[5]
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_holder = slide.shapes.title
        title_holder.text = "注意事项"
        
        tip_points = [f"{i+1}. {tip}" for i, tip in enumerate(itinerary['tips'])]
        self._add_bullet_points(slide, tip_points, 1, 2, 14, 6)
    
    def generate_packing_slide(self, itinerary: Dict[str, Any]):
        slide_layout = self.prs.slide_layouts[5]
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_holder = slide.shapes.title
        title_holder.text = "出行清单"
        
        packing_points = [f"• {item}" for item in itinerary['packing_list']]
        self._add_bullet_points(slide, packing_points, 1, 2, 14, 6)
    
    def generate_document(self, itinerary: Dict[str, Any], output_path: str):
        self.generate_cover(itinerary)
        self.generate_overview(itinerary)
        self.generate_daily_slides(itinerary)
        self.generate_food_slide(itinerary)
        self.generate_budget_slide(itinerary)
        self.generate_tips_slide(itinerary)
        self.generate_packing_slide(itinerary)
        self.prs.save(output_path)
        return output_path